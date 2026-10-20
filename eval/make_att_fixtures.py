"""Build the attachment test set: fixture files in eval/att/, eval/att_corpus.json and eval/att_queries.json.

Dev tool. Run it once on the Mac; the outputs are committed. It needs openpyxl, python-docx, python-pptx and Pillow
(the DSH bundled Python has them), a TrueType font with Turkish letters, and the LibreOffice kit for the PDF files.
Usage: <python> eval/make_att_fixtures.py --node <node> --cli <libreoffice-kit cli.js> [--font <ttf>]

Each target mail says nothing about its content in the body; the words of the query are only in the attachment.
Image and scanned-PDF targets need OCR (Windows only); on the Mac they are expected misses.
"""
import argparse
import json
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
OUT = HERE / "att"

TEXTS = {
    "xlsx_sep": [["Ürün kodu", "Ürün adı", "Miktar", "Birim"], ["RF-110", "Galvanizli raf dikmesi", 37, "adet"],
                 ["TP-020", "Paletli transpalet", 4, "adet"], ["KS-300", "Koruyucu köşebent", 250, "adet"]],
    "xlsx_oct": [["Ürün kodu", "Ürün adı", "Miktar", "Birim"], ["PK-120", "Plastik kasa", 120, "adet"],
                 ["SF-040", "Streç film", 40, "rulo"], ["EP-008", "Euro palet", 64, "adet"]],
    "docx": ["KİRA SÖZLEŞMESİ TASLAĞI", "Madde 4: Kira bedeli her yıl TÜFE oranında artırılır.",
             "Madde 5: Kiralanan depo alanı 2.400 metrekaredir ve rampa sayısı altıdır.",
             "Madde 9: Fesih bildirimi üç ay önceden yazılı olarak yapılır."],
    "pptx": [("Üçüncü Çeyrek Değerlendirmesi", "İade oranı yüzde 2,7 seviyesine indi"),
             ("Müşteri memnuniyeti", "Anket puanı 4,6 / 5 oldu; şikâyetlerin yarısı teslimat saatiyle ilgili")],
    "pdf_a": ["FATURA", "Fatura No: TR4471", "Hizmet: Soğuk zincir taşıma hizmeti",
              "Tutar: 18.750 TL (KDV hariç)", "Vade: 30 gün"],
    "pdf_b": ["FATURA", "Fatura No: TR4472", "Hizmet: Kuru yük taşıma hizmeti", "Tutar: 9.300 TL (KDV hariç)",
              "Vade: 45 gün"],
    "scan": ["YANGIN TÜPÜ DOLUM VE KONTROL RAPORU", "Kontrol tarihi: 14 Ağustos 2026",
             "Kontrol edilen tüp sayısı: 18", "Sonraki kontrol: Şubat 2027"],
    "jpg": ["Konteyner mühür numarası 558213", "Kapı menteşesi kırık, zemin ıslak"],
    "png": ["Sipariş ekranı", "Hata kodu E-4012: stok rezervasyonu başarısız"],
    "txt": "Kalibrasyon cihazı seri numarası KX-2291.\r\nBir sonraki kalibrasyon mart ayında yapılacak.\r\n",
    "csv": "Ürün;Fiyat\r\nStreç film;2.150 TL\r\nKoli bandı;85 TL\r\nBalonlu naylon;640 TL\r\n",
    "zip_docx": ["Bakım notu", "Nem ölçer sensörü arızalı, değişim için teklif istendi.",
                "Soğuk oda kapısının contası da kontrol edilecek."],
    "html": ("<html><head><meta charset='utf-8'><title>İç bülten</title><style>p{color:#333}</style></head><body>"
             "<h1>İç Bülten – Kasım</h1><p>Yeni vardiya düzeni 3 Kasım pazartesi başlıyor.</p>"
             "<p>Gece vardiyası saat 23.00&#39;te başlayacak; servis güzergâhı değişmedi.</p></body></html>"),
}


def image(lines, path, font_path, size=(1500, 420), font_size=46):
    from PIL import Image, ImageDraw, ImageFont
    im = Image.new("RGB", size, "white")
    d = ImageDraw.Draw(im)
    f = ImageFont.truetype(font_path, font_size)
    y = 40
    for line in lines:
        d.text((50, y), line, fill="black", font=f)
        y += int(font_size * 1.6)
    if path.suffix == ".jpg":
        im.save(path, quality=90)
    else:
        im.save(path)


def docx(paras, path, picture=None):
    from docx import Document
    from docx.shared import Cm
    doc = Document()
    for i, p in enumerate(paras):
        doc.add_heading(p, level=1) if i == 0 else doc.add_paragraph(p)
    if picture:
        doc.add_picture(str(picture), width=Cm(17))
    doc.save(path)


def to_pdf(src, dst, node, cli):
    subprocess.run([node, cli, "convert", "--input", str(src), "--output", str(dst)], check=True,
                   capture_output=True, timeout=300)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--node", required=True)
    ap.add_argument("--cli", required=True)
    ap.add_argument("--font", default="/System/Library/Fonts/Supplemental/Arial Unicode.ttf")
    a = ap.parse_args()
    from openpyxl import Workbook
    from pptx import Presentation

    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir()
    tmp = Path(tempfile.mkdtemp())
    for key, name in (("xlsx_sep", "Stok_Sayim_Eylul.xlsx"), ("xlsx_oct", "Stok_Sayim_Ekim.xlsx")):
        wb = Workbook()
        ws = wb.active
        ws.title = "Sayım"
        for row in TEXTS[key]:
            ws.append(row)
        wb.save(OUT / name)
    docx(TEXTS["docx"], OUT / "Kira_Sozlesmesi_Taslak.docx")
    prs = Presentation()
    for title, body in TEXTS["pptx"]:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        s.placeholders[1].text = body
    prs.save(OUT / "Ceyrek_Sunumu_Q3.pptx")
    for key, name in (("pdf_a", "Fatura_TR4471.pdf"), ("pdf_b", "Fatura_TR4472.pdf")):
        docx(TEXTS[key], tmp / (key + ".docx"))
        to_pdf(tmp / (key + ".docx"), OUT / name, a.node, a.cli)
    # Scanned PDF: one page that holds only a picture of the text (no text layer).
    image(TEXTS["scan"], tmp / "scan.png", a.font, size=(1500, 520))
    docx([""], tmp / "scan.docx", picture=tmp / "scan.png")
    to_pdf(tmp / "scan.docx", OUT / "Tarama_0042.pdf", a.node, a.cli)
    image(TEXTS["jpg"], OUT / "Hasarli_Konteyner.jpg", a.font)
    image(TEXTS["png"], OUT / "Ekran_Goruntusu.png", a.font)
    (OUT / "Kalibrasyon_Notlari.txt").write_bytes(TEXTS["txt"].encode("cp1254"))
    (OUT / "Ambalaj_Fiyatlari.csv").write_bytes(TEXTS["csv"].encode("utf-8-sig"))
    docx(TEXTS["zip_docx"], tmp / "bakim_notu.docx")
    with zipfile.ZipFile(OUT / "Bakim_Arsivi.zip", "w", zipfile.ZIP_DEFLATED) as z:
        z.write(tmp / "bakim_notu.docx", "Bakim/bakim_notu.docx")
        z.writestr("Bakim/okuma.txt", "Arşiv dosyası.")
    (OUT / "Ic_Bulten.html").write_text(TEXTS["html"], encoding="utf-8")
    shutil.rmtree(tmp)

    def m(i, subject, body, date, sender, atts=()):
        return {"id": i, "folder": "F-INBOX", "cls": "IPM.Note", "mod": date, "subject": subject, "from_name": sender,
                "from_addr": sender.split()[0].lower() + "@ornek.com.tr", "to": [["Furkan Çakır", "furkan@ornek.com.tr"]],
                "cc": [], "date": date, "attachments": list(atts), "att_paths": ["att/" + x for x in atts],
                "body": body, "lang": "tr"}

    ek = "Merhaba,\n\nDosya ektedir. Bilginize.\n\nİyi çalışmalar."
    msgs = [
        m("A01", "Sayım dosyası", ek, "2026-09-30T17:10:00", "Kemal Arslan", ["Stok_Sayim_Eylul.xlsx"]),
        m("A02", "Sayım dosyası", ek, "2026-10-31T16:40:00", "Kemal Arslan", ["Stok_Sayim_Ekim.xlsx"]),
        m("A03", "Taslak", "Merhaba,\n\nTaslak ektedir, yorumlarınızı bekliyorum.", "2026-06-12T10:05:00",
          "Selin Koç", ["Kira_Sozlesmesi_Taslak.docx"]),
        m("A04", "Sözleşme imzalandı", "Merhaba,\n\nSözleşme dün imzalandı, asıl nüsha muhasebede.",
          "2026-07-02T09:30:00", "Selin Koç"),
        m("A05", "Sunum", ek, "2026-10-08T14:20:00", "Deniz Aydın", ["Ceyrek_Sunumu_Q3.pptx"]),
        m("A06", "Fatura", ek, "2026-08-14T11:00:00", "Hızlı Lojistik", ["Fatura_TR4471.pdf"]),
        m("A07", "Fatura", ek, "2026-08-21T11:00:00", "Hızlı Lojistik", ["Fatura_TR4472.pdf"]),
        m("A08", "Tarama", "", "2026-08-15T08:45:00", "Tarayıcı", ["Tarama_0042.pdf"]),
        m("A09", "Fotoğraf", "Merhaba,\n\nSahadan gelen fotoğraf ektedir.", "2026-09-03T13:15:00", "Murat Şahin",
          ["Hasarli_Konteyner.jpg"]),
        m("A10", "Bakar mısın?", "Selam,\n\nEkteki ekrana bakabilir misin?", "2026-10-02T15:50:00", "Ece Yıldız",
          ["Ekran_Goruntusu.png"]),
        m("A11", "Notlar", ek, "2026-05-20T10:10:00", "Okan Er", ["Kalibrasyon_Notlari.txt"]),
        m("A12", "Liste", ek, "2026-04-07T09:00:00", "Ayla Tunç", ["Ambalaj_Fiyatlari.csv"]),
        m("A13", "Arşiv", ek, "2026-03-18T16:30:00", "Okan Er", ["Bakim_Arsivi.zip"]),
        m("A14", "Bülten", "Merhaba,\n\nBülten ektedir.", "2026-10-28T08:00:00", "İnsan Kaynakları", ["Ic_Bulten.html"]),
        m("A20", "Toplantı saati", "Yarınki toplantı 14.00'e alındı.", "2026-10-09T09:00:00", "Deniz Aydın"),
        m("A21", "Rampa bakımı", "Üç numaralı rampanın bakımı cuma günü yapılacak.", "2026-09-11T10:00:00", "Okan Er"),
        m("A22", "Servis", "Servis aracı pazartesi sabah gecikebilir.", "2026-10-30T07:30:00", "İnsan Kaynakları"),
        m("A23", "Fiyat teklifi", "Koli fiyatlarında yüzde 5 artış olacak.", "2026-04-01T11:00:00", "Ayla Tunç"),
    ]
    corpus = {"now": "2026-11-05T12:00:00", "store": {"id": "STORE-ATT-TEST", "name": "Ek testi"},
              "folders": [{"id": "F-INBOX", "path": "\\Gelen Kutusu"}], "messages": msgs}
    (HERE / "att_corpus.json").write_text(json.dumps(corpus, ensure_ascii=False, indent=1), encoding="utf-8")

    def q(i, text, targets, cat):
        return {"id": i, "q": text, "targets": targets, "category": cat, "lang": "tr"}

    queries = [
        q("K01", "galvanizli raf dikmesi kaç adet sayılmıştı", ["A01"], "ek-tablo"),
        q("K02", "plastik kasa ve euro palet sayımı", ["A02"], "ek-tablo"),
        q("K03", "depo alanı 2400 metrekare olan kira sözleşmesi", ["A03"], "ek-belge"),
        q("K04", "iade oranının yüzde 2,7 olduğu sunum", ["A05"], "ek-belge"),
        q("K05", "soğuk zincir taşıma faturası", ["A06"], "ek-pdf"),
        q("K06", "kuru yük taşıma faturası 9.300 TL", ["A07"], "ek-pdf"),
        q("K07", "yangın tüpü kontrol raporu", ["A08"], "ek-ocr"),
        q("K08", "konteyner mühür numarası 558213", ["A09"], "ek-ocr"),
        q("K09", "E-4012 hata kodu stok rezervasyonu", ["A10"], "ek-ocr"),
        q("K10", "KX-2291 seri numaralı kalibrasyon cihazı", ["A11"], "ek-metin"),
        q("K11", "balonlu naylon fiyatı", ["A12"], "ek-metin"),
        q("K12", "nem ölçer sensörü arızası", ["A13"], "ek-arsiv"),
        q("K13", "gece vardiyası saat 23 bülteni", ["A14"], "ek-metin"),
        q("K14", "uzay mekiği fırlatma rampası", [], "negatif"),
    ]
    (HERE / "att_queries.json").write_text(json.dumps(queries, ensure_ascii=False, indent=1), encoding="utf-8")
    for p in sorted(OUT.iterdir()):
        print(p.name, p.stat().st_size)


if __name__ == "__main__":
    main()
